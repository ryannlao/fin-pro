"""Usuários do escritório: login por pessoa, quem lançou/pagou, comprovante e histórico de cada boleto."""

from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import acesso, admin, contexto, main, usuarios
from app.arquivos import RepositorioArquivos
from app.fila_gravacao import FilaGravacao
from app.planilha import PlanilhaBoletos
from app.repositorio import RepositorioBoletos

PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


@pytest.fixture
def sistema(tmp_path, monkeypatch):
    planilha = PlanilhaBoletos(tmp_path / "boletos.xlsx", tmp_path / "backups")
    fila = FilaGravacao(planilha, tmp_path / "fila.json")
    gerenciador = usuarios.GerenciadorUsuarios(tmp_path / "usuarios.json")
    original = contexto.usuarios
    monkeypatch.setattr(contexto, "planilha", planilha)
    monkeypatch.setattr(contexto, "fila", fila)
    monkeypatch.setattr(contexto, "repositorio", RepositorioBoletos(planilha, fila))
    monkeypatch.setattr(contexto, "repositorio_pix", RepositorioBoletos(planilha, fila, "pix"))
    monkeypatch.setattr(contexto, "arquivos", RepositorioArquivos(tmp_path / "arquivos"))
    monkeypatch.setattr(contexto, "administrador", admin.GerenciadorAdmin(tmp_path / "admin.json"))
    monkeypatch.setattr(contexto, "usuarios", gerenciador)
    acesso.configurar_usuarios(gerenciador)
    yield planilha, gerenciador
    acesso.configurar_usuarios(original)


def _boleto(cliente, beneficiario="Papelaria Central"):
    resposta = cliente.post("/api/boletos", json={
        "beneficiario": beneficiario, "valor": 89.9, "vencimento": (date.today() + timedelta(days=5)).isoformat(),
        "ignorar_validacao": True})
    assert resposta.status_code == 201, resposta.text
    return resposta.json()


def test_login_por_pessoa_registra_quem_lancou_pagou_e_anexou(sistema):
    planilha, gerenciador = sistema
    with TestClient(main.app, client=("127.0.0.1", 50000)) as computador:
        # Sem usuários cadastrados: funciona como antes, e o computador aparece como autor
        assert _boleto(computador)["lancado_por"] == "Computador principal"

        # Só o administrador cadastra usuários
        assert computador.post("/api/usuarios", json={"nome": "Maria", "senha": "123456"}).status_code == 403
        computador.post("/api/admin/senha", json={"senha": "senha-do-dono", "confirmacao": "senha-do-dono"})
        assert computador.post("/api/usuarios", json={"nome": "Maria", "senha": "123"}).status_code == 422
        criado = computador.post("/api/usuarios", json={"nome": "Maria Souza", "login": "Maria", "senha": "maria-2026"})
        assert criado.status_code == 201 and criado.json()["criado"]["login"] == "maria"
        assert computador.post("/api/usuarios", json={"nome": "Outra Maria", "login": "maria", "senha": "123456"}).status_code == 422
        id_maria = criado.json()["criado"]["id"]
        assert "senha_hash" not in str(computador.get("/api/usuarios").json())

    with TestClient(main.app, client=("192.168.0.50", 50000)) as celular:
        estado = celular.get("/api/acesso/estado").json()
        assert estado["usuarios"] and not estado["autorizado"]
        assert celular.get("/api/boletos").json()["detail"]["codigo"] == "login_necessario"
        assert celular.post("/api/usuarios/entrar", json={"login": "maria", "senha": "errada"}).status_code == 401
        assert celular.post("/api/usuarios/entrar", json={"login": "ninguem", "senha": "errada"}).status_code == 401
        entrou = celular.post("/api/usuarios/entrar", json={"login": " Maria ", "senha": "maria-2026"})
        assert entrou.status_code == 200 and entrou.json()["usuario"]["nome"] == "Maria Souza"
        assert celular.get("/api/acesso/estado").json()["usuario"]["nome"] == "Maria Souza"

        boleto = _boleto(celular, "Energia Elétrica")
        assert boleto["lancado_por"] == "Maria Souza" and not boleto["pago_por"]
        pago = celular.patch(f"/api/boletos/{boleto['id']}/status", json={"status": "Pago"}).json()
        assert pago["pago_por"] == "Maria Souza"

        # Comprovante é opcional, aceita PDF ou foto e pode ser visto e removido
        assert celular.post(f"/api/boletos/{boleto['id']}/comprovante",
                            files={"arquivo": ("nota.txt", b"texto qualquer", "text/plain")}).status_code == 422
        anexado = celular.post(f"/api/boletos/{boleto['id']}/comprovante",
                               files={"arquivo": ("comprovante.pdf", PDF, "application/pdf")})
        assert anexado.status_code == 200 and anexado.json()["comprovante_nome"] == "comprovante.pdf"
        assert celular.get(f"/api/boletos/{boleto['id']}/comprovante").headers["content-type"] == "application/pdf"

        eventos = celular.get(f"/api/boletos/{boleto['id']}/historico").json()["eventos"]
        assert [(e["acao"], e["usuario"]) for e in eventos] == [
            ("Comprovante", "Maria Souza"), ("Pagamento", "Maria Souza"), ("Lançamento", "Maria Souza")]
        assert eventos[0]["detalhes"] == "Comprovante anexado: comprovante.pdf"

        # Voltar para pendente limpa o "Pago por"
        assert celular.patch(f"/api/boletos/{boleto['id']}/status", json={"status": "Pendente"}).json()["pago_por"] in ("", None)
        assert celular.delete(f"/api/boletos/{boleto['id']}/comprovante").status_code == 204
        assert celular.get(f"/api/boletos/{boleto['id']}/comprovante").status_code == 404

    wb = load_workbook(planilha.caminho)
    historico = wb["Histórico"]
    assert historico.cell(row=1, column=8).value == "Usuário"
    assert {historico.cell(row=r, column=8).value for r in range(2, historico.max_row + 1)} >= {"Computador principal", "Maria Souza"}
    cabecalho = [c.value for c in wb["Boletos"][1]]
    assert {"Lançado por", "Pago por", "Comprovante"} <= set(cabecalho)

    # Desativar a pessoa desconecta os aparelhos dela
    gerenciador.alterar(id_maria, ativo=False)
    assert not gerenciador.em_uso
    assert gerenciador.usuario_da_sessao(entrou.cookies.get(usuarios.COOKIE)) is None


def test_exigir_login_tambem_no_computador(sistema):
    _, gerenciador = sistema
    gerenciador.criar("Fábio", "fabio", "senha-fabio")
    with TestClient(main.app, client=("127.0.0.1", 50000)) as computador:
        assert computador.get("/api/boletos").status_code == 200  # padrão: o computador não pede login
        gerenciador.definir_exigir_login_no_computador(True)
        assert computador.get("/api/boletos").json()["detail"]["codigo"] == "login_necessario"
        assert computador.get("/api/admin").status_code == 200  # a área de administrador continua acessível
        assert computador.post("/api/usuarios/entrar", json={"login": "fabio", "senha": "senha-fabio"}).status_code == 200
        assert computador.get("/api/boletos").status_code == 200
        assert computador.post("/api/usuarios/minha-senha", json={"senha_atual": "errada", "senha": "nova-senha",
                                                                   "confirmacao": "nova-senha"}).status_code == 422
        assert computador.post("/api/usuarios/minha-senha", json={"senha_atual": "senha-fabio", "senha": "nova-senha",
                                                                   "confirmacao": "nova-senha"}).status_code == 200
        computador.post("/api/usuarios/sair")
        assert computador.get("/api/boletos").status_code == 401


def test_bloqueio_apos_senhas_erradas(tmp_path):
    gerenciador = usuarios.GerenciadorUsuarios(tmp_path / "usuarios.json")
    gerenciador.criar("Ana", "ana", "senha-ana")
    for _ in range(usuarios.MAX_TENTATIVAS):
        assert gerenciador.entrar("ana", "errada", "10.0.0.9")[0] is None
    token, _, bloqueio = gerenciador.entrar("ana", "senha-ana", "10.0.0.9")
    assert token is None and bloqueio > 0
    token, usuario, _ = usuarios.GerenciadorUsuarios(tmp_path / "usuarios.json").entrar("ana", "senha-ana", "10.0.0.9")
    assert token and usuario["nome"] == "Ana"
    assert "senha-ana" not in (tmp_path / "usuarios.json").read_text(encoding="utf-8")


def test_bloqueio_de_um_login_nao_tranca_os_outros_no_mesmo_endereco(tmp_path):
    """Revisão de 19/10/2026: na VPS todos chegam pelo IP do proxy; errar a senha de um não bloqueia o escritório."""
    gerenciador = usuarios.GerenciadorUsuarios(tmp_path / "usuarios.json")
    gerenciador.criar("Ana", "ana", "senha-ana")
    gerenciador.criar("Bia", "bia", "senha-bia")
    for _ in range(usuarios.MAX_TENTATIVAS):
        gerenciador.entrar("ana", "errada", "172.18.0.2")
    assert gerenciador.entrar("ana", "senha-ana", "172.18.0.2")[2] > 0
    assert gerenciador.entrar("bia", "senha-bia", "172.18.0.2")[0]
