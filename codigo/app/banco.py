"""Banco de dados (SQLite) dos boletos, Pix, centros de custo e histórico — com a planilha Excel sempre atualizada.

Por que: a planilha era o banco de dados, e cada alteração reescrevia o arquivo Excel inteiro (lento conforme cresce,
e travado quando alguém deixava o arquivo aberto no Excel). Agora:

* os dados ficam em ``dados/boletos.db`` (SQLite: rápido, gravações seguras, várias pessoas ao mesmo tempo);
* a planilha ``boletos.xlsx`` continua existindo, no mesmo formato, e é **gerada automaticamente** poucos segundos
  depois de cada alteração. Se estiver aberta no Excel, é atualizada assim que o arquivo for fechado — os
  lançamentos nunca ficam esperando;
* **edições feitas no Excel** (ex.: mudar o status ou o centro de custo de um boleto e salvar) são trazidas para o
  sistema. Só os campos que a pessoa mudou no Excel são aplicados, então nada que o sistema gravou nesse meio-tempo
  (um pagamento pelo bot, por exemplo) é desfeito. Para lançar ou excluir, use o sistema;
* na primeira vez, tudo o que está na planilha é importado para o banco, com os mesmos IDs e o histórico.

A classe tem a mesma interface de ``PlanilhaBoletos`` (listar, obter, aplicar_lote, histórico, centros...).
"""

from __future__ import annotations

import json
import logging
import os
import sqlite3
import threading
import time
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from .planilha import (
    ABA_CENTROS, ABA_HISTORICO, ABAS, BACKUPS_MANTIDOS, CAMPOS_DATA, CAMPOS_DATA_HORA, CAMPOS_MOEDA, CENTROS_PADRAO,
    TIPOS_CHAVE, PlanilhaBloqueadaError, PlanilhaBoletos, RegistroNaoEncontrado, _estilizar_cabecalho,
)

log = logging.getLogger("leitor_boletos.banco")

ESPERA_EXPORTACAO = 1.5
INTERVALO_VERIFICACAO = 15
USUARIO_EXCEL = "Planilha Excel (edição manual)"


def _para_banco(registro: dict) -> str:
    return json.dumps({k: (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in registro.items()},
                      ensure_ascii=False)


def _do_banco(texto: str) -> dict:
    registro = json.loads(texto)
    for chave in CAMPOS_DATA:
        if isinstance(registro.get(chave), str) and registro[chave]:
            registro[chave] = date.fromisoformat(registro[chave][:10])
    for chave in CAMPOS_DATA_HORA:
        if isinstance(registro.get(chave), str) and registro[chave]:
            registro[chave] = datetime.fromisoformat(registro[chave])
    return registro


def _normalizar(registro: dict, definicao) -> dict:
    """Mesmo formato que a leitura da planilha devolve (textos vazios como "", valores com 2 casas)."""
    saida = {c.chave: registro.get(c.chave) for c in definicao.colunas}
    for chave in definicao.textos:
        saida[chave] = saida.get(chave) or ""
    for chave in CAMPOS_MOEDA & set(saida):  # só os campos que a aba tem
        if isinstance(saida.get(chave), (int, float)):
            saida[chave] = round(float(saida[chave]), 2)
        elif saida.get(chave) in ("", None):
            saida[chave] = None
    if definicao.tem("status"):
        saida["status"] = saida.get("status") if saida.get("status") in definicao.status_validos else definicao.status_padrao
    return saida


def _igual(a, b) -> bool:
    if isinstance(a, float) or isinstance(b, float):
        try:
            return abs(float(a or 0) - float(b or 0)) < 0.005 and (a in (None, "")) == (b in (None, ""))
        except (TypeError, ValueError):
            return False
    return (a or None) == (b or None)


class BancoBoletos:
    def __init__(self, caminho_banco: Path, caminho_planilha: Path, pasta_backup: Path | None = None,
                 abas: dict | None = None, centros: bool = True, nome_backup: str = "boletos"):
        """``abas``: conjunto de abas (padrão: Boletos e Pix). ``centros``: usar a aba Centros de Custo e o Resumo."""
        self.caminho_banco = Path(caminho_banco)
        self.caminho = Path(caminho_planilha)  # planilha Excel (compatibilidade: exportar, abrir, backup)
        self.pasta_backup = Path(pasta_backup) if pasta_backup else None
        self.abas = abas or ABAS
        self.padrao = next(iter(self.abas))
        self.centros = centros
        self.nome_backup = nome_backup
        self.excel = PlanilhaBoletos(self.caminho, self.pasta_backup, abas=self.abas, extras=centros)
        self._lock = threading.RLock()
        self._cache: dict[str, tuple[int, list]] = {}
        self._versao = 0
        self._alterado = threading.Event()
        self._parar = threading.Event()
        self._thread: threading.Thread | None = None
        self.exportacao = {"pendente": False, "ultima": None, "erro": None, "excel_aberto": False, "ultima_importacao": None}
        self.caminho_banco.parent.mkdir(parents=True, exist_ok=True)
        self._criar_tabelas()

    # ------------------------------------------------------------------ conexão e tabelas

    @contextmanager
    def _conexao(self):
        conexao = sqlite3.connect(self.caminho_banco, timeout=30)
        try:
            yield conexao
            conexao.commit()
        except Exception:
            conexao.rollback()
            raise
        finally:
            conexao.close()

    def _criar_tabelas(self) -> None:
        with self._lock, self._conexao() as c:
            c.execute("PRAGMA journal_mode=WAL")
            c.executescript("""
                CREATE TABLE IF NOT EXISTS registros (aba TEXT NOT NULL, id INTEGER NOT NULL, dados TEXT NOT NULL,
                                                      PRIMARY KEY (aba, id));
                CREATE TABLE IF NOT EXISTS historico (seq INTEGER PRIMARY KEY AUTOINCREMENT, quando TEXT, acao TEXT,
                    registro_id INTEGER, nome TEXT, vencimento TEXT, valor REAL, detalhes TEXT, usuario TEXT);
                CREATE INDEX IF NOT EXISTS historico_registro ON historico (registro_id);
                CREATE TABLE IF NOT EXISTS centros (nome TEXT PRIMARY KEY COLLATE NOCASE, descricao TEXT, ativo INTEGER,
                    criado_em TEXT, ordem INTEGER);
                CREATE TABLE IF NOT EXISTS exportado (aba TEXT NOT NULL, id INTEGER NOT NULL, dados TEXT NOT NULL,
                                                      PRIMARY KEY (aba, id));
                CREATE TABLE IF NOT EXISTS meta (chave TEXT PRIMARY KEY, valor TEXT);
            """)

    def _meta(self, c, chave: str, valor=None, gravar: bool = False):
        if gravar:
            c.execute("INSERT OR REPLACE INTO meta (chave, valor) VALUES (?, ?)", (chave, None if valor is None else str(valor)))
            return valor
        linha = c.execute("SELECT valor FROM meta WHERE chave = ?", (chave,)).fetchone()
        return linha[0] if linha else None

    def _mudou(self) -> None:
        self._versao += 1
        self._cache = {}
        self.exportacao["pendente"] = True
        self._alterado.set()

    # ------------------------------------------------------------------ migração da planilha

    @property
    def vazio(self) -> bool:
        with self._conexao() as c:
            return c.execute("SELECT valor FROM meta WHERE chave = 'importado_em'").fetchone() is None

    def garantir(self) -> None:
        with self._lock:
            if self.vazio:
                self.migrar()
            if not self.caminho.exists():
                self.exportar()

    def migrar(self) -> bool:
        """Na primeira vez, importa a planilha existente (boletos, Pix, centros e histórico). Depois, traz edições do Excel."""
        with self._lock:
            if not self.vazio:
                return self.importar_edicoes_do_excel() > 0
            registros = {aba: [] for aba in self.abas}
            centros, historico = [], []
            if self.caminho.exists():
                try:
                    self.excel.migrar()
                except PlanilhaBloqueadaError:
                    pass  # aberta no Excel: dá para ler mesmo assim
                for aba in self.abas:
                    registros[aba] = self.excel.listar(aba)
                if self.centros:
                    centros = self.excel.listar_centros()
                historico = self._ler_historico_da_planilha()
            if not centros and self.centros:
                centros = [{"nome": n, "descricao": d, "ativo": True} for n, d in CENTROS_PADRAO]
            agora = datetime.now().replace(microsecond=0).isoformat()
            with self._conexao() as c:
                for aba, lista in registros.items():
                    c.executemany("INSERT OR REPLACE INTO registros (aba, id, dados) VALUES (?, ?, ?)",
                                  [(aba, r["id"], _para_banco(r)) for r in lista])
                c.executemany("INSERT OR IGNORE INTO centros (nome, descricao, ativo, criado_em, ordem) VALUES (?, ?, ?, ?, ?)",
                              [(ct["nome"], ct.get("descricao") or "", int(bool(ct.get("ativo", True))), agora, i)
                               for i, ct in enumerate(centros)])
                c.executemany("INSERT INTO historico (quando, acao, registro_id, nome, vencimento, valor, detalhes, usuario) "
                              "VALUES (?, ?, ?, ?, ?, ?, ?, ?)", historico)
                self._meta(c, "importado_em", agora, gravar=True)
            total = sum(len(v) for v in registros.values())
            log.info("Banco de dados criado a partir da planilha: %d lançamento(s), %d evento(s) de histórico", total, len(historico))
            self._mudou()
            self._salvar_snapshot()  # a planilha atual é o ponto de partida para detectar edições no Excel
            return True

    def _ler_historico_da_planilha(self) -> list[tuple]:
        from openpyxl import load_workbook

        wb = load_workbook(self.caminho, read_only=True, data_only=True)
        try:
            if ABA_HISTORICO not in wb.sheetnames:
                return []
            linhas = []
            for linha in wb[ABA_HISTORICO].iter_rows(min_row=2, values_only=True):
                if not linha or linha[0] is None:
                    continue
                valores = list(linha) + [None] * (8 - len(linha))
                quando, acao, id_, nome, venc, valor, detalhes, usuario = valores[:8]
                try:
                    id_ = int(id_) if id_ not in (None, "") else None
                except (TypeError, ValueError):
                    id_ = None
                linhas.append((quando.isoformat() if isinstance(quando, datetime) else str(quando),
                               str(acao or ""), id_, str(nome or ""),
                               venc.date().isoformat() if isinstance(venc, datetime) else venc.isoformat() if isinstance(venc, date) else None,
                               float(valor) if isinstance(valor, (int, float)) else None, str(detalhes or ""), str(usuario or "")))
            return linhas
        finally:
            wb.close()

    # ------------------------------------------------------------------ leitura

    def listar(self, aba: str = "") -> list[dict]:
        aba = aba or self.padrao
        definicao = self.abas[aba]
        with self._lock:
            em_cache = self._cache.get(aba)
            if not em_cache:
                with self._conexao() as c:
                    linhas = c.execute("SELECT dados FROM registros WHERE aba = ? ORDER BY id", (aba,)).fetchall()
                em_cache = (self._versao, [_normalizar(_do_banco(l[0]), definicao) for l in linhas])
                self._cache[aba] = em_cache
            return [dict(r) for r in em_cache[1]]

    def obter(self, id_: int, aba: str = "") -> dict:
        for registro in self.listar(aba):
            if registro["id"] == id_:
                return registro
        raise RegistroNaoEncontrado(id_)

    def buscar_por_codigo(self, codigo_barras: str) -> list[dict]:
        if not codigo_barras:
            return []
        return [r for r in self.listar() if r.get("codigo_barras") == codigo_barras and r.get("status") != "Cancelado"]

    def nomes_por_documento(self) -> dict[str, str]:
        import re

        mapa: dict[str, str] = {}
        for r in sorted(self.listar(), key=lambda r: r.get("atualizado_em") or r.get("data_lancamento") or datetime.min):
            documento = re.sub(r"\D", "", r.get("documento_beneficiario") or "")
            nome = (r.get("beneficiario") or "").strip()
            if len(documento) in (11, 14) and nome and r.get("status") != "Cancelado":
                mapa[documento] = nome
        return mapa

    def listar_centros(self) -> list[dict]:
        with self._conexao() as c:
            linhas = c.execute("SELECT nome, descricao, ativo FROM centros ORDER BY ordem, rowid").fetchall()
        return [{"nome": n, "descricao": d or "", "ativo": bool(a)} for n, d, a in linhas]

    def historico_do_registro(self, id_: int, aba: str = "", desde: datetime | None = None) -> list[dict]:
        sufixo = self.abas[aba or self.padrao].sufixo_acao
        outros = [d.sufixo_acao for d in self.abas.values() if d.sufixo_acao and d.sufixo_acao != sufixo]
        limite = desde.replace(second=0, microsecond=0).isoformat() if isinstance(desde, datetime) else ""
        with self._conexao() as c:
            linhas = c.execute("SELECT quando, acao, detalhes, usuario FROM historico WHERE registro_id = ? AND quando >= ? "
                               "ORDER BY seq DESC", (id_, limite)).fetchall()
        saida = []
        for quando, acao, detalhes, usuario in linhas:
            if acao == "Centro de custo":
                continue
            if sufixo and not acao.endswith(sufixo):
                continue  # evento de outra aba com o mesmo número de ID
            if not sufixo and any(acao.endswith(s) for s in outros):
                continue
            saida.append({"quando": quando[:16] if quando else None, "acao": acao.removesuffix(sufixo) if sufixo else acao,
                          "detalhes": detalhes or "", "usuario": usuario or ""})
        return saida

    def contagens(self) -> dict:
        with self._conexao() as c:
            por_aba = dict(c.execute("SELECT aba, COUNT(*) FROM registros GROUP BY aba").fetchall())
            historico = c.execute("SELECT COUNT(*) FROM historico").fetchone()[0]
            centros = c.execute("SELECT COUNT(*) FROM centros").fetchone()[0]
        return {**{aba: por_aba.get(aba, 0) for aba in self.abas}, "historico": historico, "centros": centros}

    # ------------------------------------------------------------------ escrita

    def aplicar_lote(self, operacoes: list[dict]) -> list[dict]:
        """Mesmo contrato de ``PlanilhaBoletos.aplicar_lote``, numa única transação do banco."""
        if not operacoes:
            return []
        agora = datetime.now().replace(second=0, microsecond=0)
        resultados: list[dict] = []
        with self._lock, self._conexao() as c:
            proximos: dict[str, int] = {}
            centros = {n.lower() for (n,) in c.execute("SELECT nome FROM centros")}

            def historico(acao, registro, detalhes, definicao, usuario):
                venc = registro.get(definicao.campo_data)
                valor = registro.get("valor")
                if valor is None:
                    valor = registro.get("valor_total")
                c.execute("INSERT INTO historico (quando, acao, registro_id, nome, vencimento, valor, detalhes, usuario) "
                          "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                          (agora.isoformat(), acao, registro.get("id"), registro.get(definicao.campo_nome) or "",
                           venc.isoformat() if isinstance(venc, date) else None, valor, detalhes or "", usuario or ""))

            def garantir_centro(nome) -> None:
                nome = nome.strip() if isinstance(nome, str) else ""
                if nome and self.centros and nome.lower() not in centros:
                    ordem = c.execute("SELECT COALESCE(MAX(ordem), 0) + 1 FROM centros").fetchone()[0]
                    c.execute("INSERT OR IGNORE INTO centros (nome, descricao, ativo, criado_em, ordem) VALUES (?, '', 1, ?, ?)",
                              (nome, agora.isoformat(), ordem))
                    centros.add(nome.lower())

            for operacao in operacoes:
                aba = operacao.get("aba") or self.padrao
                if aba not in self.abas:
                    resultados.append({"ok": False, "erro": "aba_invalida"})
                    continue
                definicao = self.abas[aba]
                tipo = operacao.get("tipo")
                usuario = str(operacao.get("usuario") or "")[:60]
                dados = {k: v for k, v in (operacao.get("dados") or {}).items() if k in definicao.editaveis}

                if tipo == "adicionar":
                    if aba not in proximos:
                        proximos[aba] = c.execute("SELECT COALESCE(MAX(id), 0) + 1 FROM registros WHERE aba = ?", (aba,)).fetchone()[0]
                    registro = {col.chave: None for col in definicao.colunas}
                    registro.update(dados)
                    registro.update(id=proximos[aba], data_lancamento=agora, atualizado_em=agora, lancado_por=usuario)
                    if definicao.tem("status"):
                        registro["status"] = registro.get("status") or definicao.status_padrao
                    if definicao.tem("pago_por"):
                        registro["pago_por"] = usuario if registro.get("status") == "Pago" else None
                    registro = _normalizar(registro, definicao)
                    proximos[aba] += 1
                    c.execute("INSERT INTO registros (aba, id, dados) VALUES (?, ?, ?)", (aba, registro["id"], _para_banco(registro)))
                    garantir_centro(registro.get("centro_custo"))
                    detalhe = (f"{TIPOS_CHAVE.get(registro.get('tipo_chave'), '')}: {registro.get('chave') or ''}" if aba == "pix"
                               else registro.get("linha_digitavel") or registro.get("detalhe_historico") or "")
                    historico("Lançamento" + definicao.sufixo_acao, registro, detalhe, definicao, usuario)
                    resultados.append({"ok": True, "registro": registro})
                elif tipo == "atualizar":
                    linha = c.execute("SELECT dados FROM registros WHERE aba = ? AND id = ?", (aba, operacao.get("id"))).fetchone()
                    if linha is None:
                        resultados.append({"ok": False, "erro": "nao_encontrado"})
                        continue
                    antes = _normalizar(_do_banco(linha[0]), definicao)
                    if definicao.tem("pago_por") and "status" in dados and dados["status"] != antes.get("status"):
                        dados["pago_por"] = usuario if dados["status"] == "Pago" else None
                    dados["atualizado_em"] = agora
                    depois = _normalizar({**antes, **dados}, definicao)
                    c.execute("UPDATE registros SET dados = ? WHERE aba = ? AND id = ?", (_para_banco(depois), aba, depois["id"]))
                    garantir_centro(dados.get("centro_custo"))
                    acao, detalhes = PlanilhaBoletos._descrever_alteracoes(antes, dados, definicao)
                    historico(acao + definicao.sufixo_acao, depois, detalhes, definicao, usuario)
                    resultados.append({"ok": True, "registro": depois})
                elif tipo == "excluir":
                    linha = c.execute("SELECT dados FROM registros WHERE aba = ? AND id = ?", (aba, operacao.get("id"))).fetchone()
                    if linha is None:
                        resultados.append({"ok": False, "erro": "nao_encontrado"})
                        continue
                    registro = _normalizar(_do_banco(linha[0]), definicao)
                    c.execute("DELETE FROM registros WHERE aba = ? AND id = ?", (aba, registro["id"]))
                    historico("Exclusão" + definicao.sufixo_acao, registro,
                              registro.get("linha_digitavel") or registro.get("chave") or "", definicao, usuario)
                    resultados.append({"ok": True, "registro": registro})
                else:
                    resultados.append({"ok": False, "erro": "operacao_invalida"})
            self._mudou()
        return resultados

    def adicionar(self, dados: dict, aba: str = "", usuario: str = "") -> dict:
        return self.aplicar_lote([{"tipo": "adicionar", "aba": aba, "dados": dados, "usuario": usuario}])[0]["registro"]

    def adicionar_lote(self, lista: list[dict], aba: str = "", usuario: str = "") -> list[dict]:
        return [r["registro"] for r in self.aplicar_lote(
            [{"tipo": "adicionar", "aba": aba, "dados": d, "usuario": usuario} for d in lista])]

    def atualizar(self, id_: int, alteracoes: dict, aba: str = "", usuario: str = "") -> dict:
        resultado = self.aplicar_lote([{"tipo": "atualizar", "aba": aba, "id": id_, "dados": alteracoes, "usuario": usuario}])[0]
        if not resultado["ok"]:
            raise RegistroNaoEncontrado(id_)
        return resultado["registro"]

    def excluir(self, id_: int, aba: str = "", usuario: str = "") -> None:
        if not self.aplicar_lote([{"tipo": "excluir", "aba": aba, "id": id_, "usuario": usuario}])[0]["ok"]:
            raise RegistroNaoEncontrado(id_)

    def salvar_centro(self, nome: str, descricao: str = "", ativo: bool = True, nome_anterior: str | None = None,
                      usuario: str = "") -> dict:
        nome = (nome or "").strip()
        descricao = (descricao or "").strip()
        if not nome or len(nome) > 60:
            raise ValueError("Informe o nome do centro de custo (até 60 caracteres).")
        agora = datetime.now().replace(second=0, microsecond=0).isoformat()
        with self._lock, self._conexao() as c:
            existente = c.execute("SELECT nome FROM centros WHERE nome = ?", (nome,)).fetchone()
            if nome_anterior:
                anterior = nome_anterior.strip()
                atual = c.execute("SELECT rowid FROM centros WHERE nome = ?", (anterior,)).fetchone()
                if existente and existente[0].lower() != anterior.lower():
                    raise ValueError(f'Já existe um centro de custo chamado "{nome}".')
                if atual is None:
                    ordem = c.execute("SELECT COALESCE(MAX(ordem), 0) + 1 FROM centros").fetchone()[0]
                    c.execute("INSERT INTO centros (nome, descricao, ativo, criado_em, ordem) VALUES (?, ?, ?, ?, ?)",
                              (nome, descricao, int(ativo), agora, ordem))
                else:
                    c.execute("UPDATE centros SET nome = ?, descricao = ?, ativo = ? WHERE rowid = ?", (nome, descricao, int(ativo), atual[0]))
                atualizados = 0
                if anterior != nome:
                    for aba, id_, texto in c.execute("SELECT aba, id, dados FROM registros").fetchall():
                        registro = json.loads(texto)
                        if isinstance(registro.get("centro_custo"), str) and registro["centro_custo"].strip().lower() == anterior.lower():
                            registro["centro_custo"] = nome
                            c.execute("UPDATE registros SET dados = ? WHERE aba = ? AND id = ?", (json.dumps(registro, ensure_ascii=False), aba, id_))
                            atualizados += 1
                    detalhes = f"Renomeado de {anterior} para {nome} ({atualizados} lançamento(s) atualizado(s))"
                else:
                    detalhes = "Dados atualizados" + ("" if ativo else " · desativado")
            else:
                if existente:
                    raise ValueError(f'Já existe um centro de custo chamado "{nome}".')
                ordem = c.execute("SELECT COALESCE(MAX(ordem), 0) + 1 FROM centros").fetchone()[0]
                c.execute("INSERT INTO centros (nome, descricao, ativo, criado_em, ordem) VALUES (?, ?, ?, ?, ?)",
                          (nome, descricao, int(ativo), agora, ordem))
                detalhes = "Criado"
            c.execute("INSERT INTO historico (quando, acao, registro_id, nome, vencimento, valor, detalhes, usuario) "
                      "VALUES (?, 'Centro de custo', NULL, ?, NULL, NULL, ?, ?)", (agora, nome, detalhes, usuario or ""))
            self._mudou()
        return {"nome": nome, "descricao": descricao, "ativo": ativo}

    # ------------------------------------------------------------------ planilha Excel (exportação)

    def _montar_planilha(self) -> Workbook:
        wb = Workbook()
        primeira = True
        for aba, definicao in self.abas.items():
            ws = wb.active if primeira else wb.create_sheet(definicao.nome)
            if primeira:
                ws.title = definicao.nome
                primeira = False
            self.excel._formatar_aba(ws, definicao)
            mapa = {c.chave: i for i, c in enumerate(definicao.colunas, start=1)}
            for n, registro in enumerate(self.listar(aba), start=2):
                PlanilhaBoletos._escrever(ws, n, mapa, registro, definicao)
            ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(definicao.colunas)).column_letter}{max(ws.max_row, 2)}"
        if self.centros:
            self.excel._criar_resumo(wb)
            ws = wb.create_sheet(ABA_CENTROS)
            ws.sheet_properties.tabColor = "A86200"
            for i, (titulo, largura) in enumerate([("Centro de Custo", 30), ("Descrição", 52), ("Ativo", 9), ("Criado em", 18)], start=1):
                _estilizar_cabecalho(ws.cell(row=1, column=i, value=titulo), "8A5200")
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largura
            ws.freeze_panes = "A2"
        with self._conexao() as c:
            if self.centros:
                for nome, descricao, ativo, criado in c.execute(
                        "SELECT nome, descricao, ativo, criado_em FROM centros ORDER BY ordem, rowid"):
                    ws.append([nome, descricao or None, "Sim" if ativo else "Não", datetime.fromisoformat(criado) if criado else None])
                    ws.cell(row=ws.max_row, column=4).number_format = "dd/mm/yyyy hh:mm"
            historico = PlanilhaBoletos._garantir_historico(wb)
            for quando, acao, id_, nome, venc, valor, detalhes, usuario in c.execute(
                    "SELECT quando, acao, registro_id, nome, vencimento, valor, detalhes, usuario FROM historico ORDER BY seq"):
                try:
                    momento = datetime.fromisoformat(quando)
                except (TypeError, ValueError):
                    momento = quando
                PlanilhaBoletos._registrar_historico(
                    historico, momento, acao, {"id": id_, "beneficiario": nome, "vencimento": date.fromisoformat(venc) if venc else None,
                                               "valor": valor}, detalhes, "beneficiario", usuario)
        return wb

    def exportar(self) -> bool:
        """Gera a planilha Excel a partir do banco. Devolve False se o arquivo estiver aberto no Excel (tenta de novo depois)."""
        with self._lock:
            self.importar_edicoes_do_excel()
            versao = self._versao
            wb = self._montar_planilha()
            self.caminho.parent.mkdir(parents=True, exist_ok=True)
            try:
                self.excel._verificar_gravavel() if self.caminho.exists() else None
                self.excel._salvar(wb)  # grava num temporário e substitui (com o backup diário da planilha)
            except PlanilhaBloqueadaError:
                self.exportacao.update(excel_aberto=True, pendente=True, erro=None)
                return False
            except Exception as exc:
                log.exception("Falha ao gerar a planilha Excel")
                self.exportacao.update(pendente=True, erro=str(exc))
                return False
            self._salvar_snapshot()
            self.exportacao.update(ultima=datetime.now().isoformat(timespec="seconds"), erro=None, excel_aberto=False,
                                   pendente=self._versao != versao)
            self._backup_diario_banco()
            return True

    def exportar_agora(self) -> bool:
        return self.exportar()

    def _salvar_snapshot(self) -> None:
        """Guarda o que foi escrito na planilha e a data do arquivo: base para saber o que alguém mudou no Excel."""
        with self._conexao() as c:
            c.execute("DELETE FROM exportado")
            for aba in self.abas:
                c.executemany("INSERT INTO exportado (aba, id, dados) VALUES (?, ?, ?)",
                              [(aba, r["id"], _para_banco(r)) for r in self.listar(aba)])
            if self.caminho.exists():
                self._meta(c, "planilha_mtime", self.caminho.stat().st_mtime_ns, gravar=True)

    def importar_edicoes_do_excel(self) -> int:
        """Se a planilha foi salva por alguém no Excel, aplica no banco os campos que a pessoa alterou."""
        if not self.caminho.exists():
            return 0
        with self._lock:
            with self._conexao() as c:
                guardado = self._meta(c, "planilha_mtime")
                exportados = {(aba, id_): texto for aba, id_, texto in c.execute("SELECT aba, id, dados FROM exportado")}
            mtime = self.caminho.stat().st_mtime_ns
            if guardado is None or int(guardado) == mtime:
                return 0
            try:
                leitor = PlanilhaBoletos(self.caminho, abas=self.abas, extras=self.centros)
                operacoes = []
                for aba, definicao in self.abas.items():
                    atuais = {r["id"]: r for r in self.listar(aba)}
                    for linha in leitor.listar(aba):
                        base_texto = exportados.get((aba, linha["id"]))
                        if base_texto is None or linha["id"] not in atuais:
                            continue
                        base = _normalizar(_do_banco(base_texto), definicao)
                        mudancas = {k: linha.get(k) for k in definicao.editaveis
                                    if k in linha and not _igual(linha.get(k), base.get(k)) and not _igual(linha.get(k), atuais[linha["id"]].get(k))}
                        if mudancas:
                            operacoes.append({"tipo": "atualizar", "aba": aba, "id": linha["id"], "dados": mudancas, "usuario": USUARIO_EXCEL})
            except Exception:
                log.exception("Não foi possível ler as edições feitas na planilha Excel")
                return 0
            if operacoes:
                self.aplicar_lote(operacoes)
                log.info("Planilha Excel: %d lançamento(s) alterado(s) manualmente foram trazidos para o sistema", len(operacoes))
            with self._conexao() as c:
                self._meta(c, "planilha_mtime", mtime, gravar=True)
            self.exportacao["ultima_importacao"] = datetime.now().isoformat(timespec="seconds") if operacoes else self.exportacao["ultima_importacao"]
            return len(operacoes)

    def _backup_diario_banco(self) -> None:
        if not self.pasta_backup:
            return
        destino = self.pasta_backup / f"{self.nome_backup}-{date.today():%Y-%m-%d}.db"
        if destino.exists():
            return
        self.pasta_backup.mkdir(parents=True, exist_ok=True)
        self.copiar_banco(destino)
        antigos = sorted(self.pasta_backup.glob(f"{self.nome_backup}-*.db"))
        for arquivo in antigos[:-BACKUPS_MANTIDOS]:
            arquivo.unlink(missing_ok=True)

    def copiar_banco(self, destino: Path) -> None:
        """Cópia consistente do banco (mesmo com o sistema gravando), para backup."""
        destino = Path(destino)
        temporario = destino.with_suffix(".tmp")
        with sqlite3.connect(self.caminho_banco) as origem, sqlite3.connect(temporario) as copia:
            origem.backup(copia)
        copia.close()
        os.replace(temporario, destino)

    def status(self) -> dict:
        tamanho = self.caminho_banco.stat().st_size if self.caminho_banco.exists() else 0
        return {**self.exportacao, **self.contagens(), "banco": str(self.caminho_banco), "tamanho_banco": tamanho,
                "planilha": str(self.caminho)}

    # ------------------------------------------------------------------ sincronização em segundo plano

    def iniciar_sincronizacao(self) -> None:
        if self._thread and self._thread.is_alive():
            return
        self._parar.clear()

        def laco():
            while not self._parar.is_set():
                self._alterado.wait(INTERVALO_VERIFICACAO)
                if self._parar.is_set():
                    return
                if self._alterado.is_set():
                    time.sleep(ESPERA_EXPORTACAO)  # junta alterações feitas em sequência numa única gravação
                    self._alterado.clear()
                try:
                    if self.exportacao["pendente"]:
                        self.exportar()
                    elif self.importar_edicoes_do_excel():
                        self.exportar()
                except Exception:
                    log.exception("Falha na sincronização com a planilha Excel")

        self._thread = threading.Thread(target=laco, name="sincronizar-planilha", daemon=True)
        self._thread.start()

    def parar(self) -> None:
        self._parar.set()
        self._alterado.set()
        if self._thread:
            self._thread.join(timeout=10)
        if self.exportacao["pendente"]:
            try:
                self.exportar()
            except Exception:
                log.exception("Falha ao gerar a planilha ao encerrar")
